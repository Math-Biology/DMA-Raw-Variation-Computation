# LINKED-TO: [REQ-COMP-P26.0073]
import os
import pytest
import numpy as np
import pandas as pd
from openpyxl import Workbook

from raw_variation import clean_filename, load_config, compute_percentage_variations, run, run_exported


CONFIG_TEMPLATE = """<config>
    <file_config>
        <input_path>{input_path}</input_path>
        <output_path>{output_path}</output_path>
        <exported_visits_file>{exported_visits_file}</exported_visits_file>
    </file_config>
    <visit_config>
        <threshold_1>200</threshold_1>
    </visit_config>
    <it_config>
        <original>original</original>
        <percentage_variation>percentage_variation</percentage_variation>
        <new_base>new_base</new_base>
        <fane_values>fane_values</fane_values>
        <visual_correction>visual_correction</visual_correction>
        <bmri>BMRi</bmri>
        <final_report>final_report</final_report>
    </it_config>
</config>"""


def _build_exported_excel(path):
    """Build a minimal exported_visits-format Excel file for testing.

    Visita_42: two markers across two measurement points. Values are stored as
               strings, matching the real file format. The global minimum for
               PT-1 is 50 (Marker_B) and for PT-2 is 100 (Marker_B), so:
                 Marker_A percentage variation = 500 %, Marker_B = 0 %.
    Visita_99: only two Excel rows (header + metadata), df.shape[0] < 3;
               must be skipped without raising an error.
    """
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Visita_42")
    ws1.append(["Patient ID", "First Name", "Last Name", "Gender",
                "Age at Visit", "Date of Birth", "Visit ID", "Visit Date"])
    ws1.append([1001, "Test", "User", "F", None, "1990-01-01", 42, "2024-03-15"])
    ws1.append(["Marker", "PT - 1 [1]", "PT - 2 [2]", None, None, None, None, None])
    ws1.append(["Base",          "100",       "200",   None, None, None, None, None])
    ws1.append(["Marker_A [10]", "300",       "600",   None, None, None, None, None])
    ws1.append(["Marker_B [20]",  "50",       "100",   None, None, None, None, None])

    ws2 = wb.create_sheet("Visita_99")
    ws2.append(["Patient ID", "First Name", "Last Name", "Gender",
                "Age at Visit", "Date of Birth", "Visit ID", "Visit Date"])
    ws2.append([2002, "Short", "Visit", "M", None, "1985-06-01", 99, "2024-04-01"])

    wb.save(path)


# ---------------------------------------------------------------------------
# clean_filename
# ---------------------------------------------------------------------------

class TestCleanFilename:
    def test_forbidden_chars_replaced(self):
        assert clean_filename('a<b>c:d"e/f\\g|h?i*j.k') == 'a_b_c_d_e_f_g_h_i_j_k'

    def test_clean_string_unchanged(self):
        assert clean_filename('Protocol_A1') == 'Protocol_A1'

    def test_custom_replacement_char(self):
        assert clean_filename('a/b', replacement='-') == 'a-b'

    def test_empty_string(self):
        assert clean_filename('') == ''


# ---------------------------------------------------------------------------
# load_config
# ---------------------------------------------------------------------------

class TestLoadConfig:
    SAMPLE_CONFIG = CONFIG_TEMPLATE.format(
        input_path="./Input/",
        output_path="./Output/",
        exported_visits_file="exported_visits.xlsx",
    )

    def test_loads_file_config(self, tmp_path):
        (tmp_path / "config.xml").write_text(self.SAMPLE_CONFIG)
        file_config, _, _ = load_config(str(tmp_path / "config.xml"))
        assert file_config["input_path"] == "./Input/"
        assert file_config["output_path"] == "./Output/"

    def test_loads_exported_visits_file_key(self, tmp_path):
        (tmp_path / "config.xml").write_text(self.SAMPLE_CONFIG)
        file_config, _, _ = load_config(str(tmp_path / "config.xml"))
        assert file_config["exported_visits_file"] == "exported_visits.xlsx"

    def test_visit_config_values_are_floats(self, tmp_path):
        (tmp_path / "config.xml").write_text(self.SAMPLE_CONFIG)
        _, visit_config, _ = load_config(str(tmp_path / "config.xml"))
        assert visit_config["threshold_1"] == 200.0
        assert isinstance(visit_config["threshold_1"], float)

    def test_loads_it_config_suffixes(self, tmp_path):
        (tmp_path / "config.xml").write_text(self.SAMPLE_CONFIG)
        _, _, it_config = load_config(str(tmp_path / "config.xml"))
        assert it_config["original"] == "original"
        assert it_config["percentage_variation"] == "percentage_variation"

    def test_missing_file_raises(self):
        with pytest.raises(FileNotFoundError):
            load_config("/nonexistent/path/config.xml")


# ---------------------------------------------------------------------------
# compute_percentage_variations
# ---------------------------------------------------------------------------

class TestComputePercentageVariations:
    def test_basic_variation(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [100.0,       150.0,    200.0],
        })
        result = compute_percentage_variations(df)
        assert result["param_A"].iloc[0] == pytest.approx(50.0)
        assert result["param_A"].iloc[1] == pytest.approx(100.0)

    def test_baseline_is_global_minimum_not_row_0(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [100.0,       200.0,    50.0],
        })
        result = compute_percentage_variations(df)
        assert result["param_A"].iloc[0] == pytest.approx(300.0)
        assert result["param_A"].iloc[1] == pytest.approx(0.0)

    def test_row_0_dropped_from_output(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [100.0,       150.0,    200.0],
        })
        result = compute_percentage_variations(df)
        assert len(result) == 2

    def test_label_column_not_processed(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [100.0,       150.0,    200.0],
        })
        result = compute_percentage_variations(df)
        assert list(result["label"]) == ["time_1", "time_2"]

    def test_column_with_no_positive_values_skipped(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1"],
            "param_A": [np.nan,      np.nan],
            "param_B": [100.0,       200.0],
        })
        result = compute_percentage_variations(df)
        assert pd.isna(result["param_A"].iloc[0])
        assert result["param_B"].iloc[0] == pytest.approx(100.0)

    def test_nan_in_data_propagates_to_output(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [100.0,       np.nan,   200.0],
        })
        result = compute_percentage_variations(df)
        assert pd.isna(result["param_A"].iloc[0])
        assert result["param_A"].iloc[1] == pytest.approx(100.0)

    def test_multiple_columns_computed_independently(self):
        df = pd.DataFrame({
            "label":   ["baseline", "time_1"],
            "param_A": [50.0,        100.0],
            "param_B": [200.0,       400.0],
        })
        result = compute_percentage_variations(df)
        assert result["param_A"].iloc[0] == pytest.approx(100.0)
        assert result["param_B"].iloc[0] == pytest.approx(100.0)


# ---------------------------------------------------------------------------
# run — legacy folder-based integration
# ---------------------------------------------------------------------------

class TestRun:
    VISIT = "VISIT01"
    SHEET = "Protocol"
    FILENAME = "test_data"

    @pytest.fixture
    def workspace(self, tmp_path):
        visit_dir = tmp_path / "Input" / self.VISIT
        visit_dir.mkdir(parents=True)
        output_dir = tmp_path / "Output"
        output_dir.mkdir()

        df = pd.DataFrame({
            "label":   ["baseline", "time_1", "time_2"],
            "param_A": [0.0,         200.0,    50.0],
        })
        df.to_excel(str(visit_dir / f"{self.FILENAME}.xlsx"), index=False, sheet_name=self.SHEET)

        config_path = tmp_path / "config.xml"
        config_path.write_text(CONFIG_TEMPLATE.format(
            input_path=str(tmp_path / "Input") + os.sep,
            output_path=str(output_dir) + os.sep,
            exported_visits_file="exported_visits.xlsx",
        ))

        return {
            "config_path": str(config_path),
            "output_dir": output_dir,
            "visit_dir": visit_dir,
        }

    def _output_folder(self, workspace):
        return workspace["output_dir"] / f"{self.VISIT}_{self.FILENAME}_{self.SHEET}"

    def test_both_output_files_created(self, workspace):
        run(config_path=workspace["config_path"])
        folder = self._output_folder(workspace)
        assert (folder / f"{self.FILENAME}_{self.SHEET}_original.xlsx").exists()
        assert (folder / f"{self.FILENAME}_{self.SHEET}_percentage_variation.xlsx").exists()

    def test_percentage_variation_values_correct(self, workspace):
        run(config_path=workspace["config_path"])
        folder = self._output_folder(workspace)
        result = pd.read_excel(folder / f"{self.FILENAME}_{self.SHEET}_percentage_variation.xlsx")
        assert result["param_A"].iloc[0] == pytest.approx(300.0)
        assert result["param_A"].iloc[1] == pytest.approx(0.0)

    def test_original_file_contains_all_rows(self, workspace):
        run(config_path=workspace["config_path"])
        folder = self._output_folder(workspace)
        original = pd.read_excel(folder / f"{self.FILENAME}_{self.SHEET}_original.xlsx")
        assert len(original) == 3

    def test_variation_file_excludes_baseline_row(self, workspace):
        run(config_path=workspace["config_path"])
        folder = self._output_folder(workspace)
        result = pd.read_excel(folder / f"{self.FILENAME}_{self.SHEET}_percentage_variation.xlsx")
        assert len(result) == 2

    def test_excluded_sheets_produce_no_output(self, workspace):
        excel_path = workspace["visit_dir"] / f"{self.FILENAME}.xlsx"
        with pd.ExcelWriter(str(excel_path), engine="openpyxl") as writer:
            pd.DataFrame({"label": ["baseline", "time_1"], "val": [100.0, 200.0]}).to_excel(
                writer, sheet_name=self.SHEET, index=False)
            pd.DataFrame({"note": ["x"]}).to_excel(writer, sheet_name="Notes", index=False)
            pd.DataFrame({"note": ["x"]}).to_excel(writer, sheet_name="Screening Protocol", index=False)

        run(config_path=workspace["config_path"])

        output_dir = workspace["output_dir"]
        assert not (output_dir / f"{self.VISIT}_{self.FILENAME}_Notes").exists()
        assert not (output_dir / f"{self.VISIT}_{self.FILENAME}_Screening_Protocol").exists()


# ---------------------------------------------------------------------------
# run_exported — multi-sheet Excel integration
# ---------------------------------------------------------------------------

class TestRunExported:
    EXPORTED_FILE = "exported_visits.xlsx"

    @pytest.fixture(scope="class")
    def workspace(self, tmp_path_factory):
        tmp_path = tmp_path_factory.mktemp("exported")
        input_dir = tmp_path / "Input"
        input_dir.mkdir()
        output_dir = tmp_path / "Output"
        output_dir.mkdir()

        _build_exported_excel(str(input_dir / self.EXPORTED_FILE))

        config_path = tmp_path / "config.xml"
        config_path.write_text(CONFIG_TEMPLATE.format(
            input_path=str(input_dir) + os.sep,
            output_path=str(output_dir) + os.sep,
            exported_visits_file=self.EXPORTED_FILE,
        ))

        run_exported(config_path=str(config_path))

        return {"output_dir": output_dir}

    def test_original_csv_created(self, workspace):
        assert (workspace["output_dir"] / "all_visits_original.csv").exists()

    def test_variation_csv_created(self, workspace):
        assert (workspace["output_dir"] / "all_visits_percentage_variation.csv").exists()

    def test_original_csv_has_required_columns(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_original.csv")
        assert list(df.columns) == ["visit_id", "patient_id", "visit_date", "marker", "point", "value"]

    def test_variation_csv_has_required_columns(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_percentage_variation.csv")
        assert list(df.columns) == ["visit_id", "patient_id", "visit_date", "marker", "point", "percentage_variation"]

    def test_metadata_columns_populated(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_original.csv")
        assert (df["visit_id"] == 42).all()
        assert (df["patient_id"] == 1001).all()

    def test_original_csv_includes_base_row(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_original.csv")
        assert "Base" in df["marker"].values

    def test_variation_csv_excludes_base_row(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_percentage_variation.csv")
        assert "Base" not in df["marker"].values

    def test_percentage_variation_values_correct(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_percentage_variation.csv")
        row_a = df[(df["marker"] == "Marker_A [10]") & (df["point"] == "PT - 1 [1]")]
        row_b = df[(df["marker"] == "Marker_B [20]") & (df["point"] == "PT - 1 [1]")]
        assert row_a["percentage_variation"].iloc[0] == pytest.approx(500.0)
        assert row_b["percentage_variation"].iloc[0] == pytest.approx(0.0)

    def test_string_values_converted_to_numeric(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_original.csv")
        assert pd.api.types.is_numeric_dtype(df["value"])

    def test_short_sheet_skipped_without_error(self, workspace):
        df = pd.read_csv(workspace["output_dir"] / "all_visits_original.csv")
        assert 99 not in df["visit_id"].values


    def test_all_sheets_invalid_exits_without_crash(self, tmp_path_factory):
        tmp_path = tmp_path_factory.mktemp("all_invalid")
        input_dir = tmp_path / "Input"
        input_dir.mkdir()
        output_dir = tmp_path / "Output"
        output_dir.mkdir()

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Visita_1")
        ws.append(["Patient ID", "First Name", "Last Name", "Gender",
                   "Age at Visit", "Date of Birth", "Visit ID", "Visit Date"])
        ws.append([999, "Test", "Solo", "F", None, "1990-01-01", 1, "2024-01-01"])
        wb.save(str(input_dir / self.EXPORTED_FILE))

        config_path = tmp_path / "config.xml"
        config_path.write_text(CONFIG_TEMPLATE.format(
            input_path=str(input_dir) + os.sep,
            output_path=str(output_dir) + os.sep,
            exported_visits_file=self.EXPORTED_FILE,
        ))

        run_exported(config_path=str(config_path))
        assert not (output_dir / "all_visits_original.csv").exists()
        assert not (output_dir / "all_visits_percentage_variation.csv").exists()
